import io
import base64
import openpyxl
import logging

from odoo import models, fields, api, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class ImportPrecioWizard(models.TransientModel):
    _name = 'import.precio.wizard'
    _description = 'Importar Precios desde Excel'

    file_xls = fields.Binary('Archivo Excel', required=True)
    file_xls_name = fields.Char('Nombre Archivo')

    def action_import(self):
        self.ensure_one()
        if not self.file_xls:
            raise UserError(_('Seleccione un archivo Excel.'))

        file_data = base64.b64decode(self.file_xls)
        try:
            wb = openpyxl.load_workbook(
                filename=io.BytesIO(file_data),
                data_only=True
            )
        except Exception:
            raise UserError(_(
                'El archivo debe ser un Excel valido (.xlsx).'
            ))

        ws = wb.active
        actualizados = 0
        no_encontrados_list = []
        errores = []

        for row in ws.iter_rows(
            min_row=2, max_row=ws.max_row, values_only=True
        ):
            codigo = row[0]
            descripcion = row[1]
            precio = row[2]

            if not codigo:
                continue

            codigo = str(codigo).strip()
            if not codigo:
                continue

            try:
                precio_val = float(precio) if precio else 0.0
            except (ValueError, TypeError):
                errores.append(
                    'Codigo %s: precio invalido "%s"' % (codigo, precio)
                )
                continue

            if precio_val < 0:
                errores.append(
                    'Codigo %s: precio debe ser 0 o mayor a 0' % codigo
                )
                continue

            product = self.env['product.product'].search([
                ('default_code', '=', codigo)
            ], limit=1)

            if product:
                product.list_price = precio_val
                actualizados += 1
            else:
                product_tmpl = self.env['product.template'].search([
                    ('default_code', '=', codigo)
                ], limit=1)
                if product_tmpl:
                    product_tmpl.list_price = precio_val
                    actualizados += 1
                else:
                    no_encontrados_list.append(codigo)

        mensaje = 'Productos actualizados: %d' % actualizados
        if no_encontrados_list:
            max_show = 30
            no_show = no_encontrados_list[:max_show]
            mensaje += '\nNo encontrados: %d' % len(no_encontrados_list)
            mensaje += '\nLista: %s' % ', '.join(no_show)
            if len(no_encontrados_list) > max_show:
                mensaje += '\n... y %d mas' % (len(no_encontrados_list) - max_show)

        if errores:
            _logger.warning(
                'Errores en importacion de precios:\n%s' % '\n'.join(errores)
            )

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Importacion Completada'),
                'message': mensaje,
                'type': 'success' if not (no_encontrados_list or errores) else 'warning',
                'sticky': True,
            },
        }
